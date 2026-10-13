"""Los Programas de Estudio de Ajedrez de 1.° a 11.°, con la estructura de los
programas del MEP, y el planeamiento de las 200 lecciones de cada año, en Word
con la marca de agua de Oscar Angulo Cubero. Ver «Los programas de estudio de
Ajedrez» en docs/decisiones/cursos-y-material.md.

El contenido está en datos/grado-N.json (datos/GUIA.md dice cómo se escribe y
datos/validar.py lo revisa). Los .docx no se editan a mano: se cambia el JSON
y se corre esto.

Uso: python3 herramientas/programas-estudio/generar.py   (necesita python-docx)
     -> material/programas-de-estudio/*.docx, el .zip con todos y
        js/programas-estudio-lista.js"""
import json, math, os, sys
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm, RGBColor
import docbase as d
import comun as c

AQUI = os.path.dirname(os.path.abspath(__file__))
DATOS = os.environ.get("DATOS", os.path.join(AQUI, "datos"))
SALIDA = os.environ.get("SALIDA", os.path.join(AQUI, "..", "..", "material", "programas-de-estudio"))
AÑO = "2027"


NOMBRE_PARTE = {1: "Programa de Estudio de Ajedrez, parte 1 (fundamentos)",
                2: "Programa de Estudio de Ajedrez, parte 2 (malla curricular)"}


def nombre_archivo(g, tipo):
    return f"{g:02d} {c.ORDINAL[g]} - {tipo}.docx"


def portada(doc, g, datos, titulo, subtitulo):
    for _ in range(5):
        doc.add_paragraph()
    d.parrafo(doc, titulo.upper(), tam=24, negrita=True, alin="c", color=d.AZUL, despues=6)
    d.parrafo(doc, subtitulo, tam=14, alin="c", color=d.GRIS, despues=30)
    d.parrafo(doc, c.ORDINAL[g], tam=30, negrita=True, alin="c", color=d.AZUL, despues=4)
    d.parrafo(doc, datos["ciclo"], tam=13, alin="c", despues=4)
    d.parrafo(doc, f"Edad de referencia: {datos['edad']}", tam=11, alin="c", color=d.GRIS, despues=60)
    d.parrafo(doc, "Autor", tam=10, alin="c", color=d.GRIS, despues=0)
    d.parrafo(doc, d.AUTOR, tam=16, negrita=True, alin="c", despues=40)
    d.parrafo(doc, f"Costa Rica, {AÑO}", tam=11, alin="c", despues=4)
    d.salto(doc)
    p = d.parrafo(doc, f"© {AÑO} {d.AUTOR}. Todos los derechos reservados. Se permite su uso en el aula citando al autor. "
                       "Prohibida su venta o reproducción con fines comerciales sin autorización escrita.",
                  tam=8.5, cursiva=True, alin="c", color=d.GRIS)


def seccion_horizontal(doc, horizontal):
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    ancho, alto = Cm(21.59), Cm(27.94)
    if horizontal:
        s.orientation = WD_ORIENT.LANDSCAPE
        s.page_width, s.page_height = alto, ancho
    else:
        s.orientation = WD_ORIENT.PORTRAIT
        s.page_width, s.page_height = ancho, alto
    return s


def programa(g, datos, parte):
    ciclo = c.CICLO_DE[g]
    titulo = f"Programa de Estudio de Ajedrez. {c.ORDINAL[g]}"
    doc = d.documento(titulo, f"Programa de estudio de Ajedrez, {c.ORDINAL[g]}, {datos['ciclo']}",
                      f"ajedrez; programa de estudio; {c.ORDINAL[g].lower()}; {datos['ciclo']}; Costa Rica")
    portada(doc, g, datos, "Programa de Estudio de Ajedrez",
            "Parte 1 de 2: Fundamentos y orientaciones" if parte == 1 else "Parte 2 de 2: Malla curricular y evaluación")

    doc.add_heading("Contenido", 1)
    partes = ["1. Presentación", "2. Fundamentación", "3. Enfoque curricular y mediación pedagógica",
              "4. Las habilidades y el ajedrez", "5. Ejes transversales y valores",
              f"6. Perfil de salida de {c.ORDINAL[g].lower()}", "7. Orientaciones metodológicas",
              "8. Orientaciones para la evaluación", "9. Distribución del tiempo", "10. Malla curricular",
              "11. Proyectos", "12. Instrumentos de evaluación", "13. Glosario", "14. Referencias"]
    d.parrafo(doc, "El programa se presenta en dos archivos. La parte 1 contiene las secciones 1 a 9 y la parte 2, "
                   "las secciones 10 a 14.", tam=9.5, cursiva=True, color=d.GRIS)
    for i, p in enumerate(partes):
        d.parrafo(doc, p + ("   (parte 1)" if i < 9 else "   (parte 2)"), despues=2,
                  negrita=(i < 9) == (parte == 1))

    if parte == 1:
        d.salto(doc)
        doc.add_heading("1. Presentación", 1)
        for p in c.PRESENTACION:
            d.parrafo(doc, p.format(ordinal_min=c.ORDINAL[g].lower(), ciclo=datos["ciclo"], del_ciclo=("de la " if datos["ciclo"].startswith("Educación") else "del ") + datos["ciclo"]), alin="j")
        doc.add_heading(f"Sobre {c.ORDINAL[g].lower()}", 2)
        for p in datos["presentacion"]:
            d.parrafo(doc, p, alin="j")

        doc.add_heading("2. Fundamentación", 1)
        for p in c.FUNDAMENTACION:
            d.parrafo(doc, p, alin="j")

        doc.add_heading("3. Enfoque curricular y mediación pedagógica", 1)
        for p in c.ENFOQUE:
            d.parrafo(doc, p, alin="j")
        d.tabla(doc, ["Momento", "Qué sucede"], [[a, b] for a, b in c.MOMENTOS], [4.5, 12.5])

        doc.add_heading("4. Las habilidades y el ajedrez", 1)
        d.parrafo(doc, "El marco curricular organiza las habilidades en cuatro dimensiones. La tabla muestra cómo "
                       "las trabaja el ajedrez. En la malla curricular, cada aprendizaje esperado se vincula con una "
                       "de ellas y con un indicador observable.", alin="j")
        d.tabla(doc, ["Dimensión", "Habilidad", "Cómo la desarrolla el ajedrez"], [list(h) for h in c.HABILIDADES], [3.6, 3.8, 9.6])

        doc.add_heading("5. Ejes transversales y valores", 1)
        d.tabla(doc, ["Eje transversal", "Aporte desde el ajedrez"], [list(e) for e in c.EJES], [5.5, 11.5])
        d.parrafo(doc, "")
        d.parrafo(doc, "Los valores que se cultivan en todas las lecciones son el respeto, la responsabilidad, la "
                       "honestidad, la solidaridad, la perseverancia, la tolerancia y el juego limpio. Cada unidad "
                       "destaca uno o dos de ellos.", alin="j")

        doc.add_heading(f"6. Perfil de salida de {c.ORDINAL[g].lower()}", 1)
        d.parrafo(doc, "Conocimientos que se esperan al iniciar el año:", negrita=True)
        d.vinetas(doc, datos["conocimientos_previos"])
        d.parrafo(doc, "Al finalizar el año, la persona estudiante:", negrita=True)
        d.vinetas(doc, [quitar_inicio(p) for p in datos["perfil_salida"]])

        doc.add_heading("7. Orientaciones metodológicas", 1)
        for p in c.ORIENT_METOD[ciclo] + c.ORIENT_METOD["comun"]:
            d.parrafo(doc, p, alin="j")

        doc.add_heading("8. Orientaciones para la evaluación", 1)
        for p in c.EVAL_COMUN:
            d.parrafo(doc, p, alin="j")
        d.parrafo(doc, "Componentes de la calificación", negrita=True)
        d.tabla(doc, ["Componente", "Porcentaje"], [list(x) for x in c.COMPONENTES[ciclo]], [8, 4])
        d.parrafo(doc, c.NOTA_COMPONENTES, tam=9, cursiva=True, color=d.GRIS)

        doc.add_heading("9. Distribución del tiempo", 1)
        d.parrafo(doc, "El año se planifica en 40 semanas de 5 lecciones de 40 minutos, para un total de 200 "
                       "lecciones, distribuidas en dos semestres. Los días marcados como flexibles en el "
                       "planeamiento permiten ajustar el plan al calendario escolar oficial del año.", alin="j")
        filas = []
        for u in datos["unidades"]:
            a, b = u["dias"]
            filas.append([f"{u['n']}. {u['titulo']}", "I" if b <= 100 else "II", f"{a} a {b}",
                          f"{math.ceil(a/5)} a {math.ceil(b/5)}", str(b - a + 1)])
        d.tabla(doc, ["Unidad", "Semestre", "Lecciones", "Semanas", "Total"], filas, [8, 2, 2.4, 2.4, 2.2])

    if parte == 2:
        # malla en horizontal
        seccion_horizontal(doc, True)
        doc.add_heading("10. Malla curricular", 1)
        for u in datos["unidades"]:
            doc.add_heading(f"Unidad {u['n']}. {u['titulo']}", 2)
            a, b = u["dias"]
            d.parrafo(doc, f"**Tiempo probable:** {b - a + 1} lecciones (de la {a} a la {b}).   **Eje transversal:** {u['eje']}.   "
                           f"**Valores:** {', '.join(u['valores'])}.", tam=9.5)
            d.parrafo(doc, f"**Propósito:** {u['proposito']}", tam=9.5, alin="j")
            filas = []
            for i, ae in enumerate(u["aprendizajes"], 1):
                co = ae["contenidos"]
                cont = ([f"**Conceptuales:** {'; '.join(co['conceptuales'])}"] if co.get("conceptuales") else []) + \
                       ([f"**Procedimentales:** {'; '.join(co['procedimentales'])}"] if co.get("procedimentales") else []) + \
                       ([f"**Actitudinales:** {'; '.join(co['actitudinales'])}"] if co.get("actitudinales") else [])
                niv = [f"**Inicial:** {ae['niveles']['inicial']}", f"**Intermedio:** {ae['niveles']['intermedio']}",
                       f"**Avanzado:** {ae['niveles']['avanzado']}"]
                filas.append([f"**{u['n']}.{i}** {ae['ae']}", [f"**{ae['habilidad']}**", ae["indicador_habilidad"]],
                              cont, ae["estrategias"], niv])
            d.tabla(doc, ["Aprendizaje esperado", "Habilidad e indicador", "Contenidos", "Estrategias de mediación sugeridas",
                          "Indicadores del aprendizaje esperado (niveles de desempeño)"],
                    filas, [4.2, 3.6, 5.6, 5.2, 5.4], tam=8.5)

        seccion_horizontal(doc, False)
        doc.add_heading("11. Proyectos", 1)
        for pr in datos["proyectos"]:
            doc.add_heading(f"{'I' if pr['semestre'] == 1 else 'II'} semestre: {pr['titulo']}", 2)
            d.parrafo(doc, pr["descripcion"], alin="j")
            d.parrafo(doc, f"**Producto:** {pr['producto']}")
            if pr.get("avances"):
                d.tabla(doc, ["Lección", "Avance"], [[str(a["dia"]), a["que"]] for a in pr["avances"]], [2.5, 14.5])
            d.parrafo(doc, "Criterios de valoración:", negrita=True)
            d.vinetas(doc, pr["criterios"])

        doc.add_heading("12. Instrumentos de evaluación", 1)
        d.parrafo(doc, "Instrumentos listos para usar en las evaluaciones sumativas del año. La persona docente los "
                       "puede adaptar a su grupo manteniendo los criterios.", alin="j")
        for ins in datos["instrumentos"]:
            doc.add_heading(ins["nombre"], 3)
            if ins.get("criterios"):
                d.tabla(doc, ["Criterio", "Inicial (1)", "Intermedio (2)", "Avanzado (3)"],
                        [[x["criterio"], x["inicial"], x["intermedio"], x["avanzado"]] for x in ins["criterios"]],
                        [4.4, 4.2, 4.2, 4.2], tam=8.5)
            if ins.get("items"):
                d.tabla(doc, ["Indicador", "Sí", "No", "Observaciones"], [[x, "", "", ""] for x in ins["items"]],
                        [10, 1.2, 1.2, 4.6], tam=9)
            d.parrafo(doc, "")

        doc.add_heading("13. Glosario", 1)
        for t, s in c.GLOSARIO:
            d.parrafo(doc, f"**{t}.** {s}", alin="j", despues=3)

        doc.add_heading("14. Referencias", 1)
        for r in c.REFERENCIAS:
            p = d.parrafo(doc, r, despues=3)
            p.paragraph_format.left_indent = Cm(0.8)
            p.paragraph_format.first_line_indent = Cm(-0.8)

    d.encabezado_y_pie(doc, f"{d.AUTOR}  ·  Programa de Estudio de Ajedrez  ·  {c.ORDINAL[g]}  ·  Parte {parte}",
                       f"Ajedrez · {c.ORDINAL[g]}")
    ruta = os.path.join(SALIDA, nombre_archivo(g, NOMBRE_PARTE[parte]))
    d.guardar(doc, ruta)
    return ruta


def quitar_inicio(p):
    for pref in ("Al finalizar", "Al terminar"):
        if p.startswith(pref) and "estudiante" in p[:80]:
            resto = p.split("estudiante", 1)[1].lstrip(" ,:")
            return resto[0].upper() + resto[1:]
    return p


def planeamiento(g, datos, sem=None, unidad=None):
    if unidad:
        return planeamiento_unidad(g, datos, unidad)
    rom = {1: "I", 2: "II"}.get(sem, "")
    titulo = f"Planeamiento diario de Ajedrez. {c.ORDINAL[g]}" + (f". {rom} semestre" if sem else "")
    doc = d.documento(titulo, f"Planeamiento de las 200 lecciones de Ajedrez, {c.ORDINAL[g]}",
                      f"ajedrez; planeamiento didáctico; lecciones; {c.ORDINAL[g].lower()}; Costa Rica")
    portada(doc, g, datos, "Planeamiento diario de Ajedrez",
            f"{rom} semestre: lecciones {1 if sem == 1 else 101} a {100 if sem == 1 else 200}" if sem else "Las 200 lecciones del curso lectivo")

    doc.add_heading("Cómo usar este planeamiento", 1)
    for p in [
        f"Este planeamiento desarrolla, lección por lección, el Programa de Estudio de Ajedrez de "
        f"{c.ORDINAL[g].lower()}. Son 200 lecciones de 40 minutos organizadas en 40 semanas y dos semestres. "
        "Cada lección indica la unidad, el aprendizaje esperado al que apunta (con su número en la malla "
        "curricular), el objetivo, los cuatro momentos de la mediación, la evaluación y los materiales.",
        "Los tiempos de cada momento son de referencia: focalización 5 minutos, exploración 15, contrastación 10 y "
        "aplicación 10. Si el grupo necesita más tiempo en un momento, se recorta otro, nunca la aplicación.",
        "La casilla de fecha queda en blanco para que cada docente la complete con el calendario escolar del año. "
        "Las lecciones marcadas como **flexibles** son de refuerzo o ampliación: si por feriados, actos o cambios "
        "de calendario hay menos días lectivos, son las primeras que se pueden omitir sin romper la secuencia.",
        "Las referencias a «Ajedrez en el aula» y a los Manuales 1 a 5 remiten a los materiales de consulta "
        "indicados en el programa. Las partidas y ejercicios citados se montan en el tablero mural o se "
        "entregan impresos.",
    ]:
        d.parrafo(doc, p, alin="j")

    doc.add_heading("Cronograma por semanas", 1)
    unid = {u["n"]: u for u in datos["unidades"]}
    filas = []
    for s in (range(1, 21) if sem == 1 else range(21, 41) if sem == 2 else range(1, 41)):
        dias = datos["dias"][(s - 1) * 5: s * 5]
        us = sorted({x["u"] for x in dias})
        filas.append([str(s), f"{dias[0]['d']} a {dias[-1]['d']}", ", ".join(str(x) for x in us),
                      [x["tema"] for x in dias]])
    d.tabla(doc, ["Semana", "Lecciones", "Unidad", "Temas"], filas, [1.6, 2.2, 1.6, 11.6], tam=8)

    doc.add_heading("Unidades del año", 2)
    for u in [u for u in datos["unidades"] if not sem or (u["dias"][1] <= 100) == (sem == 1)]:
        a, b = u["dias"]
        d.parrafo(doc, f"**Unidad {u['n']}. {u['titulo']}** (lecciones {a} a {b})", despues=1)
        for i, ae in enumerate(u["aprendizajes"], 1):
            p = d.parrafo(doc, f"{u['n']}.{i}  {ae['ae']}", tam=9, despues=1)
            p.paragraph_format.left_indent = Cm(0.6)

    sem_actual = None
    u_actual = None
    for x in datos["dias"]:
        sx = 1 if x["d"] <= 100 else 2
        if sem and sx != sem:
            continue
        if sx != sem_actual:
            sem_actual = sx
            d.salto(doc)
            doc.add_heading(f"{'I' if sx == 1 else 'II'} semestre", 1)
        if x["u"] != u_actual:
            u_actual = x["u"]
            u = unid[u_actual]
            doc.add_heading(f"Unidad {u['n']}. {u['titulo']}", 2)
        leccion(doc, x, unid)

    d.encabezado_y_pie(doc, f"{d.AUTOR}  ·  Planeamiento diario de Ajedrez  ·  {c.ORDINAL[g]}" + (f"  ·  {rom} semestre" if sem else ""),
                       f"Ajedrez · {c.ORDINAL[g]}")
    ruta = os.path.join(SALIDA, nombre_archivo(g, "Planeamiento diario de Ajedrez" + (f", {rom} semestre" if sem else "")))
    d.guardar(doc, ruta)
    return ruta


def planeamiento_unidad(g, datos, n):
    u = next(u for u in datos["unidades"] if u["n"] == n)
    a, b = u["dias"]
    unid = {v["n"]: v for v in datos["unidades"]}
    titulo = f"Planeamiento diario de Ajedrez. {c.ORDINAL[g]}. Unidad {n}"
    doc = d.documento(titulo, f"Planeamiento de las lecciones {a} a {b} de Ajedrez, {c.ORDINAL[g]}",
                      f"ajedrez; planeamiento didáctico; lecciones; {c.ORDINAL[g].lower()}; Costa Rica")
    for _ in range(4):
        doc.add_paragraph()
    d.parrafo(doc, "PLANEAMIENTO DIARIO DE AJEDREZ", tam=22, negrita=True, alin="c", color=d.AZUL, despues=4)
    d.parrafo(doc, f"{c.ORDINAL[g]}  ·  {datos['ciclo']}", tam=13, alin="c", color=d.GRIS, despues=24)
    d.parrafo(doc, f"Unidad {n}", tam=28, negrita=True, alin="c", color=d.AZUL, despues=4)
    d.parrafo(doc, u["titulo"], tam=16, alin="c", despues=6)
    d.parrafo(doc, f"Lecciones {a} a {b}  ·  semanas {math.ceil(a/5)} a {math.ceil(b/5)}  ·  "
                   f"{'I' if b <= 100 else 'II'} semestre", tam=11, alin="c", color=d.GRIS, despues=50)
    d.parrafo(doc, "Autor", tam=10, alin="c", color=d.GRIS, despues=0)
    d.parrafo(doc, d.AUTOR, tam=16, negrita=True, alin="c", despues=30)
    d.parrafo(doc, f"Costa Rica, {AÑO}", tam=11, alin="c", despues=30)
    d.parrafo(doc, f"© {AÑO} {d.AUTOR}. Todos los derechos reservados. Se permite su uso en el aula citando al autor. "
                   "Prohibida su venta o reproducción con fines comerciales sin autorización escrita.",
              tam=8.5, cursiva=True, alin="c", color=d.GRIS)
    d.salto(doc)
    doc.add_heading(f"Unidad {n}. {u['titulo']}", 1)
    d.parrafo(doc, f"**Propósito.** {u['proposito']}", alin="j")
    d.parrafo(doc, f"**Eje transversal:** {u['eje']}.   **Valores:** {', '.join(u['valores'])}.")
    d.parrafo(doc, "Aprendizajes esperados de la unidad:", negrita=True, despues=2)
    for i, ae in enumerate(u["aprendizajes"], 1):
        p = d.parrafo(doc, f"**{n}.{i}**  {ae['ae']}", despues=2)
        p.paragraph_format.left_indent = Cm(0.6)
    d.parrafo(doc, "Cada lección dura 40 minutos y sigue los cuatro momentos de la mediación: focalización (5 min), "
                   "exploración (15), contrastación (10) y aplicación (10). La fecha se completa con el calendario "
                   "escolar del año. Las lecciones **flexibles** son de refuerzo o ampliación y se pueden omitir si "
                   "el calendario trae menos días lectivos.", tam=9.5, alin="j", color=d.GRIS)
    doc.add_heading("Cronograma de la unidad", 2)
    dias = [x for x in datos["dias"] if a <= x["d"] <= b]
    filas = []
    for s in range(math.ceil(a / 5), math.ceil(b / 5) + 1):
        ds = [x for x in dias if math.ceil(x["d"] / 5) == s]
        filas.append([str(s), f"{ds[0]['d']} a {ds[-1]['d']}", [f"{x['d']}. {x['tema']}" + (" (flexible)" if x.get("flex") else "") for x in ds]])
    d.tabla(doc, ["Semana", "Lecciones", "Temas"], filas, [1.8, 2.4, 12.8], tam=8.5)
    d.salto(doc)
    doc.add_heading("Lecciones", 1)
    for x in dias:
        leccion(doc, x, unid)
    d.encabezado_y_pie(doc, f"{d.AUTOR}  ·  Planeamiento diario de Ajedrez  ·  {c.ORDINAL[g]}  ·  Unidad {n}",
                       f"Ajedrez · {c.ORDINAL[g]} · Unidad {n}")
    ruta = os.path.join(SALIDA, nombre_archivo(g, f"Planeamiento Unidad {n} (lecciones {a} a {b})"))
    d.guardar(doc, ruta)
    return ruta


def leccion(doc, x, unid):
    sem = math.ceil(x["d"] / 5)
    dia_sem = (x["d"] - 1) % 5 + 1
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    t.autofit = False
    cab = t.add_row().cells
    cab[0].merge(cab[1])
    marca = "   ·   Lección flexible" if x.get("flex") else ""
    d.texto_celda(cab[0], f"Lección {x['d']}   ·   Semana {sem}, día {dia_sem}{marca}   ·   Fecha: ____________",
                  negrita=True, tam=9.5, color="FFFFFF", junto=True)
    d.sombrear(cab[0], "1F3A5F")
    filas = [
        ("Tema", f"**{x['tema']}**"),
        ("Aprendizaje esperado", f"{x['ae']}. {ae_texto(unid, x)}"),
        ("Objetivo", x["objetivo"]),
        ("Focalización (5 min)", x["foc"]),
        ("Exploración (15 min)", x["exp"]),
        ("Contrastación (10 min)", x["con"]),
        ("Aplicación (10 min)", x["apl"]),
        ("Evaluación", f"**{x['eval']['tipo']}.** {x['eval']['como']}"),
        ("Materiales", x["mat"]),
    ]
    if x.get("tarea"):
        filas.append(("Tarea", x["tarea"]))
    for et, v in filas:
        r = t.add_row().cells
        d.texto_celda(r[0], et, negrita=True, tam=8.5, color="1F3A5F", junto=True)
        d.texto_celda(r[1], v, tam=9, junto=True)
    for fila in t.rows:
        d.no_partir(fila)
    d.fijar_anchos(t, [3.4, 13.6])
    # la última fila no se pega a la ficha siguiente
    for celda in t.rows[-1].cells:
        for p in celda.paragraphs:
            p.paragraph_format.keep_with_next = False
    esp = doc.add_paragraph()
    esp.paragraph_format.space_after = Pt(2)


def ae_texto(unid, x):
    un, k = x["ae"].split(".")
    return unid[int(un)]["aprendizajes"][int(k) - 1]["ae"]


def serie(grados):
    doc = d.documento("Programas de Estudio de Ajedrez. Alcance y secuencia",
                      "Alcance y secuencia de los programas de Ajedrez de primer a undécimo año",
                      "ajedrez; programa de estudio; alcance y secuencia; Costa Rica")
    for _ in range(6):
        doc.add_paragraph()
    d.parrafo(doc, "PROGRAMAS DE ESTUDIO DE AJEDREZ", tam=24, negrita=True, alin="c", color=d.AZUL, despues=6)
    d.parrafo(doc, "Educación General Básica y Educación Diversificada", tam=14, alin="c", color=d.GRIS, despues=30)
    d.parrafo(doc, "Alcance y secuencia de primer a undécimo año", tam=18, negrita=True, alin="c", color=d.AZUL, despues=80)
    d.parrafo(doc, "Autor", tam=10, alin="c", color=d.GRIS, despues=0)
    d.parrafo(doc, d.AUTOR, tam=16, negrita=True, alin="c", despues=40)
    d.parrafo(doc, f"Costa Rica, {AÑO}", tam=11, alin="c")
    d.salto(doc)
    doc.add_heading("La serie", 1)
    for p in [
        "La serie está formada por once programas de estudio de Ajedrez, uno por cada año escolar, desde primer "
        "año de la Educación General Básica hasta undécimo año de la Educación Diversificada. Cada programa "
        "sigue la estructura de los programas de estudio del Ministerio de Educación Pública: presentación, "
        "fundamentación, enfoque, habilidades, ejes transversales, perfil de salida, orientaciones "
        "metodológicas y de evaluación, malla curricular con aprendizajes esperados, indicadores y niveles de "
        "desempeño, proyectos e instrumentos de evaluación.",
        "Cada programa va acompañado de su planeamiento diario: las 200 lecciones del curso lectivo, una por "
        "día, organizadas en los cuatro momentos de la mediación pedagógica (focalización, exploración, "
        "contrastación y aplicación), con su evaluación, materiales y tareas.",
        "Los contenidos avanzan en espiral. En el I Ciclo se aprende el tablero, las piezas, las reglas y los "
        "primeros mates, con mucho juego y movimiento. En el II Ciclo entran la notación, el reloj, la táctica "
        "y los finales básicos. En el III Ciclo se trabaja la estrategia, la evaluación de posiciones, las "
        "aperturas por sus ideas, la historia del juego y la organización de torneos. En la Educación "
        "Diversificada el estudiantado profundiza, analiza con criterio, investiga la relación del ajedrez con "
        "la matemática y la ciencia, y enseña a otras personas.",
        "Quien llega a un año sin haber cursado los anteriores encuentra en la primera unidad un repaso que le "
        "permite incorporarse. Por eso cada año puede impartirse también de forma independiente.",
    ]:
        d.parrafo(doc, p, alin="j")
    doc.add_heading("Archivos de la serie", 2)
    filas = []
    for g in grados:
        datos = json.load(open(os.path.join(DATOS, f"grado-{g}.json"), encoding="utf-8"))
        nombres = [nombre_archivo(g, NOMBRE_PARTE[1]), nombre_archivo(g, NOMBRE_PARTE[2])] + [
            nombre_archivo(g, f"Planeamiento Unidad {u['n']} (lecciones {u['dias'][0]} a {u['dias'][1]})") for u in datos["unidades"]]
        filas.append([c.ORDINAL[g], c.NOMBRE_CICLO[c.CICLO_DE[g]], [n[:-5] for n in nombres]])
    d.tabla(doc, ["Año", "Ciclo", "Archivos"], filas, [2.6, 3.8, 10.6], tam=8.5)
    seccion_horizontal(doc, True)
    doc.add_heading("Alcance y secuencia", 1)
    for g in grados:
        datos = json.load(open(os.path.join(DATOS, f"grado-{g}.json"), encoding="utf-8"))
        doc.add_heading(f"{c.ORDINAL[g]} ({datos['edad']})", 3)
        filas = [[f"{u['n']}. {u['titulo']}", "I" if u["dias"][1] <= 100 else "II",
                  [ae["ae"] for ae in u["aprendizajes"]]] for u in datos["unidades"]]
        d.tabla(doc, ["Unidad", "Semestre", "Aprendizajes esperados"], filas, [6.5, 1.8, 15.6], tam=8.5)
    d.encabezado_y_pie(doc, f"{d.AUTOR}  ·  Programas de Estudio de Ajedrez  ·  Alcance y secuencia", "Ajedrez · Alcance y secuencia")
    ruta = os.path.join(SALIDA, "00 Alcance y secuencia - Programas de Estudio de Ajedrez.docx")
    d.guardar(doc, ruta)
    return ruta


def manifiesto(grados):
    """js/programas-estudio-lista.js: la lista que pinta admin.html#programas. Se genera; no se edita a mano."""
    años = []
    for g in grados:
        datos = json.load(open(os.path.join(DATOS, f"grado-{g}.json"), encoding="utf-8"))
        archivos = [{"titulo": "Programa de estudio, parte 1 (fundamentos)", "archivo": nombre_archivo(g, NOMBRE_PARTE[1])},
                    {"titulo": "Programa de estudio, parte 2 (malla curricular)", "archivo": nombre_archivo(g, NOMBRE_PARTE[2])}]
        for u in datos["unidades"]:
            a, b = u["dias"]
            archivos.append({"titulo": f"Planeamiento, unidad {u['n']}: {u['titulo']} (lecciones {a} a {b})",
                             "archivo": nombre_archivo(g, f"Planeamiento Unidad {u['n']} (lecciones {a} a {b})")})
        años.append({"grado": g, "nombre": c.ORDINAL[g], "ciclo": datos["ciclo"], "edad": datos["edad"], "archivos": archivos})
    js = ("/* Generado por herramientas/programas-estudio/generar.py: no se edita a mano.\n"
          "   Los Programas de Estudio de Ajedrez (1.° a 11.°) que se bajan desde admin.html#programas.\n"
          "   Viven en material/programas-de-estudio/, detrás del candado del worker: solo administración. */\n"
          "window.PROGRAMAS_ESTUDIO = " + json.dumps({"carpeta": CARPETA_WEB, "serie": SERIE, "zip": ZIP, "años": años},
                                                     ensure_ascii=False, indent=1) + ";\n")
    open(os.path.join(RAIZ, "js", "programas-estudio-lista.js"), "w", encoding="utf-8").write(js)


RAIZ = os.path.normpath(os.path.join(AQUI, "..", ".."))
CARPETA_WEB = "material/programas-de-estudio/"
SERIE = "00 Alcance y secuencia - Programas de Estudio de Ajedrez.docx"
ZIP = "Programas de Estudio de Ajedrez (todos).zip"


if __name__ == "__main__":
    import zipfile
    grados = [int(x) for x in sys.argv[1:]] or list(range(1, 12))
    os.makedirs(SALIDA, exist_ok=True)
    rutas = []
    for g in grados:
        datos = json.load(open(os.path.join(DATOS, f"grado-{g}.json"), encoding="utf-8"))
        rutas += [programa(g, datos, 1), programa(g, datos, 2)] + [planeamiento_unidad(g, datos, n) for n in range(1, 7)]
    if len(grados) == 11:
        rutas.append(serie(grados))
        with zipfile.ZipFile(os.path.join(SALIDA, ZIP), "w", zipfile.ZIP_STORED) as z:
            for r in sorted(rutas):
                z.write(r, os.path.join("Programas de Estudio de Ajedrez", os.path.basename(r)))
        manifiesto(grados)
    print(len(rutas), "documentos en", SALIDA)
