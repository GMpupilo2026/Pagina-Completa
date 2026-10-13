"""Base común de los documentos: página, estilos, marca de agua, pie y datos del archivo."""
import copy, datetime, os, re, zipfile, shutil, tempfile
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
from docx.shared import Pt, Cm, RGBColor, Emu

AQUI = os.path.dirname(os.path.abspath(__file__))
MARCA = os.path.join(AQUI, "marca.png")
AUTOR = "Oscar Angulo Cubero"
AZUL = RGBColor(0x1F, 0x3A, 0x5F)
GRIS = RGBColor(0x55, 0x55, 0x55)
FUENTE = "Arial"


def documento(titulo, asunto, palabras, horizontal=False):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)  # carta
    if horizontal:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for lado in ("left_margin", "right_margin"):
        setattr(sec, lado, Cm(2.2))
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)
    sec.header_distance, sec.footer_distance = Cm(1.0), Cm(1.0)

    est = doc.styles
    n = est["Normal"]
    n.font.name = FUENTE
    n.font.size = Pt(10.5)
    n.element.rPr.rFonts.set(qn("w:eastAsia"), FUENTE)
    n.paragraph_format.space_after = Pt(4)
    n.paragraph_format.line_spacing = 1.12
    for nivel, tam in ((1, 16), (2, 13), (3, 11.5)):
        h = est[f"Heading {nivel}"]
        h.font.name = FUENTE
        h.font.size = Pt(tam)
        h.font.bold = True
        h.font.italic = False
        h.font.color.rgb = AZUL
        rpr = h.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts"); rpr.append(rf)
        for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rf.set(qn(a), FUENTE)
        for a in list(rf.attrib):
            if "Theme" in a:
                del rf.attrib[a]
        h.paragraph_format.space_before = Pt(14 if nivel == 1 else 10)
        h.paragraph_format.space_after = Pt(5)
        h.paragraph_format.keep_with_next = True
    for nombre in ("List Bullet", "List Number"):
        est[nombre].font.name = FUENTE
        est[nombre].font.size = Pt(10.5)
        est[nombre].paragraph_format.space_after = Pt(2)

    cp = doc.core_properties
    cp.author = AUTOR
    cp.last_modified_by = AUTOR
    cp.title = titulo
    cp.subject = asunto
    cp.keywords = palabras
    cp.comments = ""
    cp.category = "Programa de estudio"
    cp.created = datetime.datetime(2026, 10, 6, 9, 0, 0)
    cp.modified = datetime.datetime(2026, 10, 6, 9, 0, 0)
    cp.revision = 1
    cp.language = "es-CR"
    # idioma del texto: español de Costa Rica (corrector ortográfico)
    rpr = n.element.get_or_add_rPr()
    lang = OxmlElement("w:lang"); lang.set(qn("w:val"), "es-CR"); rpr.append(lang)
    return doc


def _anclar_marca(parrafo, ancho_cm, pagina_ancho, pagina_alto):
    run = parrafo.add_run()
    run.add_picture(MARCA, width=Cm(ancho_cm))
    inline = run._r.find(".//" + qn("wp:inline"))
    cx, cy = int(inline.find(qn("wp:extent")).get("cx")), int(inline.find(qn("wp:extent")).get("cy"))
    graphic = inline.find(qn("a:graphic"))
    docpr = inline.find(qn("wp:docPr"))
    x = int((pagina_ancho - cx) / 2)
    y = int((pagina_alto - cy) / 2)
    anchor = parse_xml(
        f'<wp:anchor {nsdecls("wp", "a", "pic", "r")} distT="0" distB="0" distL="0" distR="0" '
        f'simplePos="0" relativeHeight="0" behindDoc="1" locked="1" layoutInCell="1" allowOverlap="1">'
        f'<wp:simplePos x="0" y="0"/>'
        f'<wp:positionH relativeFrom="page"><wp:posOffset>{x}</wp:posOffset></wp:positionH>'
        f'<wp:positionV relativeFrom="page"><wp:posOffset>{y}</wp:posOffset></wp:positionV>'
        f'<wp:extent cx="{cx}" cy="{cy}"/><wp:effectExtent l="0" t="0" r="0" b="0"/>'
        f'<wp:wrapNone/><wp:docPr id="{docpr.get("id")}" name="Marca de agua" descr="Logo de {AUTOR}"/>'
        f'<wp:cNvGraphicFramePr/></wp:anchor>'
    )
    anchor.append(copy.deepcopy(graphic))
    # rotación leve, como en los libros del sitio
    xfrm = anchor.find(".//" + qn("a:xfrm"))
    xfrm.set("rot", str(-15 * 60000))
    inline.getparent().replace(inline, anchor)


def _campo(run, instr):
    for tipo, texto in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
        if tipo:
            f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), tipo); run._r.append(f)
        elif texto == instr:
            t = OxmlElement("w:instrText"); t.set(qn("xml:space"), "preserve"); t.text = f" {instr} "; run._r.append(t)
        else:
            t = OxmlElement("w:t"); t.text = texto; run._r.append(t)


def encabezado_y_pie(doc, texto_pie, texto_encabezado=""):
    for sec in doc.sections:
        sec.different_first_page_header_footer = False
        enc = sec.header
        enc.is_linked_to_previous = False
        p = enc.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if texto_encabezado:
            r = p.add_run(texto_encabezado); r.font.size = Pt(8); r.font.color.rgb = GRIS
        _anclar_marca(p, 14.5 if sec.page_width < sec.page_height else 16, sec.page_width, sec.page_height)
        pie = sec.footer
        pie.is_linked_to_previous = False
        q = pie.paragraphs[0]
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = q.add_run(f"{texto_pie}  ·  Página "); r.font.size = Pt(8); r.font.color.rgb = GRIS
        r2 = q.add_run(); r2.font.size = Pt(8); r2.font.color.rgb = GRIS
        _campo(r2, "PAGE")


def sombrear(celda, color_hex):
    tcpr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color_hex)
    tcpr.append(shd)


def repetir_encabezado(fila):
    trpr = fila._tr.get_or_add_trPr()
    e = OxmlElement("w:tblHeader"); e.set(qn("w:val"), "true"); trpr.append(e)


def no_partir(fila):
    trpr = fila._tr.get_or_add_trPr()
    e = OxmlElement("w:cantSplit"); e.set(qn("w:val"), "true"); trpr.append(e)


def estilo_celda(doc, tam, negrita, color, junto):
    """Un estilo de párrafo por combinación: el formato va en el estilo y no repetido en cada trozo de texto."""
    nombre = f"Celda {tam} {'N' if negrita else 'R'} {color or 'X'} {'J' if junto else 'S'}".replace(".", "_")
    est = doc.styles
    try:
        return est[nombre]
    except KeyError:
        from docx.enum.style import WD_STYLE_TYPE
        e = est.add_style(nombre, WD_STYLE_TYPE.PARAGRAPH)
        e.base_style = est["Normal"]
        e.font.size = Pt(tam)
        e.font.bold = negrita
        if color: e.font.color.rgb = RGBColor.from_string(str(color))
        e.paragraph_format.space_after = Pt(2)
        e.paragraph_format.line_spacing = 1.05
        e.paragraph_format.keep_with_next = junto
        return e


def texto_celda(celda, texto, negrita=False, tam=9.5, color=None, centrado=False, junto=False):
    """Escribe en la celda; acepta lista (una viñeta por elemento) o texto con **negritas**."""
    doc = celda.part.document
    est = estilo_celda(doc, tam, negrita, color, junto)
    celda.paragraphs[0].text = ""
    items = texto if isinstance(texto, list) else [texto]
    for i, it in enumerate(items):
        p = celda.paragraphs[0] if i == 0 else celda.add_paragraph()
        p.style = est
        if centrado:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pref = "• " if isinstance(texto, list) and len(items) > 1 else ""
        escribir(p, pref + str(it))


def escribir(p, texto, tam=None, negrita=False, color=None, cursiva=False):
    partes = re.split(r"(\*\*[^*]+\*\*)", texto)
    for parte in partes:
        if not parte:
            continue
        b = parte.startswith("**") and parte.endswith("**")
        r = p.add_run(parte[2:-2] if b else parte)
        if negrita or b: r.bold = True
        if cursiva: r.italic = True
        if tam: r.font.size = Pt(tam)
        if color: r.font.color.rgb = color
    return p


def parrafo(doc, texto, tam=None, negrita=False, cursiva=False, alin=None, despues=None, color=None):
    p = doc.add_paragraph()
    escribir(p, texto, tam=tam, negrita=negrita, color=color, cursiva=cursiva)
    if alin == "c": p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if alin == "j": p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if despues is not None: p.paragraph_format.space_after = Pt(despues)
    return p


def vinetas(doc, items, estilo="List Bullet"):
    for it in items:
        p = doc.add_paragraph(style=estilo)
        escribir(p, it)


def tabla(doc, encabezados, filas, anchos_cm, color_enc="1F3A5F", tam=9, repetir=True, partir=True):
    t = doc.add_table(rows=1, cols=len(encabezados))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, e in enumerate(encabezados):
        c = t.rows[0].cells[i]
        texto_celda(c, e, negrita=True, tam=tam, color="FFFFFF")
        sombrear(c, color_enc)
    if repetir:
        repetir_encabezado(t.rows[0])
    for f in filas:
        cs = t.add_row().cells
        for i, v in enumerate(f):
            texto_celda(cs[i], v, tam=tam)
        if not partir:
            no_partir(t.rows[-1])
    fijar_anchos(t, anchos_cm)
    return t


def fijar_anchos(t, anchos_cm):
    t.autofit = False
    for i, w in enumerate(anchos_cm):
        t.columns[i].width = Cm(w)
    for fila in t.rows:
        celdas = fila._tr.tc_lst
        if len(celdas) == len(anchos_cm):
            for tc, w in zip(celdas, anchos_cm):
                tc.width = Cm(w)
        elif len(celdas) == 1:
            celdas[0].width = Cm(sum(anchos_cm))


def salto(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def guardar(doc, ruta):
    doc.save(ruta)
    limpiar_metadatos(ruta)


def limpiar_metadatos(ruta):
    """Deja en el archivo solo los datos del autor: sin rastros de la herramienta que lo armó."""
    app = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
           '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
           'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
           f'<Company>{AUTOR}</Company><DocSecurity>0</DocSecurity><ScaleCrop>false</ScaleCrop>'
           '<LinksUpToDate>false</LinksUpToDate><SharedDoc>false</SharedDoc>'
           '<HyperlinksChanged>false</HyperlinksChanged><AppVersion>16.0000</AppVersion></Properties>')
    tmp = ruta + ".tmp"
    sobran = ("word/stylesWithEffects.xml", "docProps/thumbnail.jpeg")
    with zipfile.ZipFile(ruta) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as zout:
        for item in zin.infolist():
            if item.filename in sobran or item.filename.startswith("customXml/"):
                continue
            datos = zin.read(item.filename)
            if item.filename in ("[Content_Types].xml", "_rels/.rels", "word/_rels/document.xml.rels"):
                datos = re.sub(rb'<(Override|Relationship)[^>]*(stylesWithEffects|thumbnail|customXml|itemProps)[^>]*/>', b"", datos)
            if item.filename == "word/styles.xml":
                datos = re.sub(rb"<w:latentStyles.*?</w:latentStyles>", b"", datos, flags=re.S)
                # solo los estilos que se usan
                usados = (b"Normal", b"Heading1", b"Heading2", b"Heading3", b"ListBullet", b"ListNumber",
                          b"TableGrid", b"TableNormal", b"DefaultParagraphFont", b"NoList", b"Title",
                          b"Heading1Char", b"Heading2Char", b"Heading3Char", b"TitleChar", b"Header", b"Footer",
                          b"HeaderChar", b"FooterChar")
                def filtro(m):
                    sid = re.search(rb'w:styleId="([^"]+)"', m.group(0)).group(1)
                    return m.group(0) if (sid in usados or sid.startswith(b"Celda")) else b""
                datos = re.sub(rb"<w:style [^>]*>.*?</w:style>", filtro, datos, flags=re.S)
            if item.filename == "docProps/app.xml":
                datos = app.encode()
            if item.filename == "docProps/core.xml":
                datos = re.sub(rb"<dc:description>.*?</dc:description>", b"", datos)
            if item.filename.endswith(".xml") and b"python-docx" in datos:
                datos = datos.replace(b"python-docx", AUTOR.encode())
            zout.writestr(item, datos, compress_type=zipfile.ZIP_DEFLATED, compresslevel=9)
    shutil.move(tmp, ruta)
    with zipfile.ZipFile(ruta) as z:
        for nombre in z.namelist():
            assert b"python-docx" not in z.read(nombre), nombre
