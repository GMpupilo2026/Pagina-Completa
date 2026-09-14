# -*- coding: utf-8 -*-
"""Generador de cursos de Ajedrez Integral.

Arma, a partir de un archivo de contenido (herramientas/cursos/<slug>.json),
las tres piezas que componen un curso del sitio:

  cursos/<slug>.html              portada y temario — público, sin sesión
  cursos/protegido/<slug>.html    el fragmento con las lecciones completas,
                                  que js/curso-acceso.js inyecta solo si hay
                                  sesión de Academia
  cursos/recursos/<slug>/         por lección, la presentación (.pptx) y el
                                  PDF de ejercicios para imprimir

La portada se clona de un curso existente (mismo encabezado, mismo menú, mismo
pie) para que todos los cursos sigan siendo el mismo sitio: lo único que se
reescribe es el artículo.

Cómo se corre (las dos bibliotecas no son parte del sitio, se instalan aparte):

    pip install python-pptx reportlab
    python3 herramientas/curso-generar.py estrategia-en-el-final

Después hay que agregar a mano la tarjeta del curso en cursos.html — es lo
único que no se genera, porque el orden de las tarjetas es una decisión
editorial.
"""
import json
import os
import re
import sys

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MOLDE = os.path.join(RAIZ, "cursos", "finales-practicos.html")   # de aquí sale el armazón

# Paleta del sitio (tailwind.config de cada página)
AZUL_950, AZUL_900, AZUL_800 = "0A1F33", "102A43", "243B53"
AZUL_600, AZUL_400, AZUL_300 = "334E68", "627D98", "9FB3C8"
GRIS_50 = "F0F4F8"
AMBAR_500, AMBAR_400 = "DE911D", "F0B429"
CLARO = "CADCFC"


# ---------------------------------------------------------------- utilidades
def numerar(bloques):
    """Numera las lecciones de corrido, como se ven en el temario."""
    n = 0
    for bloque in bloques:
        for leccion in bloque["lecciones"]:
            n += 1
            leccion["n"] = n
            leccion["bloque"] = bloque
    return n


def slug_de(texto):
    t = texto.lower()
    for a, b in zip("áéíóúüñ", "aeiouun"):
        t = t.replace(a, b)
    t = re.sub(r"[^a-z0-9]+", "-", t).strip("-")
    return t


def archivo_base(leccion):
    return "%02d-%s" % (leccion["n"], slug_de(leccion["titulo"]))


def escapar(t):
    return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


# ------------------------------------------------------------------ portada
def portada(curso):
    base = open(MOLDE, encoding="utf-8").read()
    total = sum(len(b["lecciones"]) for b in curso["bloques"])
    slug, titulo = curso["slug"], curso["titulo"]

    bloques_html = []
    for bloque in curso["bloques"]:
        inicio = bloque["lecciones"][0]["n"]
        items = "".join(
            '\n                        <li class="text-brand-600 dark:text-brand-300">%s'
            '<span class="block text-xs text-brand-400 dark:text-brand-500 mt-0.5">%s</span></li>'
            % (escapar(l["titulo"]), escapar(l["resumen"]))
            for l in bloque["lecciones"])
        bloques_html.append(
            '                <div>\n'
            '                    <h3 class="font-serif text-lg font-bold text-brand-800 dark:text-white mb-3">Bloque %d · %s</h3>\n'
            '                    <ol start="%d" class="space-y-2 list-decimal pl-5 marker:text-accent-600 marker:font-semibold">%s\n'
            '                    </ol>\n'
            '                </div>' % (bloque["n"], escapar(bloque["titulo"]), inicio, items))

    articulo = '''    <article class="pt-28 pb-16">
        <div class="max-w-3xl mx-auto px-4 sm:px-6 lg:px-8">
            <a href="../cursos.html" class="inline-flex items-center gap-1 text-sm text-brand-500 dark:text-brand-400 hover:text-accent-500 transition-colors mb-4"><span aria-hidden="true">←</span> Volver a Cursos</a>
            <span class="text-xs text-accent-600 font-semibold uppercase tracking-wide">{nivel}</span>
            <h1 class="font-serif text-3xl md:text-4xl font-bold text-brand-800 dark:text-white mt-2 mb-3">{titulo}</h1>
            <p class="text-brand-500 dark:text-brand-400 mb-6">{resumen}</p>
            <div aria-hidden="true" class="h-48 bg-gradient-to-br {gradiente} rounded-2xl flex items-center justify-center text-7xl mb-6">{emoji}</div>

            <p class="text-brand-600 dark:text-brand-300 mb-8">{descripcion_larga}</p>

            <div class="flex flex-wrap items-center gap-4 mb-10 bg-white dark:bg-brand-900 rounded-xl shadow-md p-5">
                <div class="flex-1 min-w-[180px]">
                    <p class="text-sm text-brand-500 dark:text-brand-400"><strong class="text-brand-800 dark:text-white">{total} lecciones</strong> · {nivel}</p>
                    <p class="text-xs text-brand-400 dark:text-brand-500 mt-0.5">Cada lección con su presentación y su PDF de ejercicios para imprimir.</p>
                </div>
                <div class="hidden items-center gap-2" id="course-access-cta" hidden>
                    <a href="../login.html?next=cursos/{slug}.html" class="bg-accent-500 hover:bg-accent-600 text-brand-900 font-semibold px-6 py-2.5 rounded-lg transition-colors whitespace-nowrap">Iniciar sesión →</a>
                </div>
            </div>

            <h2 class="font-serif text-2xl font-bold text-brand-800 dark:text-white mb-6">Temario del curso</h2>
            <div class="space-y-8">
{bloques}
            </div>

            <h2 class="font-serif text-2xl font-bold text-brand-800 dark:text-white mb-1 mt-12">Contenido completo del curso</h2>
            <p class="text-brand-500 dark:text-brand-400 text-sm mb-6">Cada lección desarrollada por completo: la idea explicada, la tarea concreta para practicarla, la presentación y el PDF de ejercicios para imprimir.</p>

            <div id="course-content-body" data-course="{slug}" class="space-y-8">
                <p class="text-sm text-brand-400 dark:text-brand-500">Cargando lecciones…</p>
            </div>

            <div class="mt-10 bg-brand-800 dark:bg-brand-900 rounded-2xl p-6 text-center hidden" id="course-login-cta" hidden>
                <p class="text-white font-serif text-lg font-bold mb-2">Contenido para alumnos de Academia</p>
                <p class="text-brand-200 text-sm mb-4">Inicia sesión con tu cuenta de Academia para ver las {total} lecciones completas de este curso, con sus recursos descargables.</p>
                <a href="../login.html?next=cursos/{slug}.html" class="inline-block bg-accent-500 hover:bg-accent-600 text-brand-900 font-semibold px-6 py-2.5 rounded-lg transition-colors">Iniciar sesión →</a>
            </div>

            <nav aria-label="Navegación entre cursos" class="mt-10 flex items-center justify-between border-t border-brand-100 dark:border-brand-800 pt-6">
                <a href="{anterior_href}" class="text-sm text-brand-500 dark:text-brand-400 hover:text-accent-500 transition-colors" aria-label="Curso anterior: {anterior_titulo}"><span aria-hidden="true">←</span> {anterior_titulo}</a>
                <a href="{siguiente_href}" class="text-sm text-brand-500 dark:text-brand-400 hover:text-accent-500 transition-colors text-right" aria-label="Siguiente curso: {siguiente_titulo}">{siguiente_titulo} <span aria-hidden="true">→</span></a>
            </nav>
        </div>
    </article>'''.format(
        nivel=escapar(curso["nivel"]), titulo=escapar(titulo), resumen=escapar(curso["resumen"]),
        gradiente=curso["gradiente"], emoji=curso["emoji"],
        descripcion_larga=escapar(curso["descripcion_larga"]),
        total=total, slug=slug, bloques="\n".join(bloques_html),
        anterior_href=curso["anterior"]["href"], anterior_titulo=escapar(curso["anterior"]["titulo"]),
        siguiente_href=curso["siguiente"]["href"], siguiente_titulo=escapar(curso["siguiente"]["titulo"]))

    salida = re.sub(r'<article class="pt-28 pb-16">.*?</article>', lambda m: articulo, base, flags=re.S)
    salida = salida.replace("<title>Finales Prácticos — Ajedrez Integral</title>",
                            "<title>%s — Ajedrez Integral</title>" % escapar(titulo))
    salida = re.sub(r'<meta name="description" content="[^"]*">',
                    '<meta name="description" content="Temario completo del curso %s: %s">'
                    % (escapar(titulo), escapar(curso["resumen"])), salida, count=1)
    return salida


# ---------------------------------------------------- fragmento con lecciones
def protegido(curso):
    slug = curso["slug"]
    boton = ('<a href="../recursos/{slug}/{arch}{suf}" class="inline-flex items-center gap-1.5 '
             'bg-brand-100 dark:bg-brand-700 text-brand-700 dark:text-brand-100 text-xs font-medium '
             'px-3 py-1.5 rounded-lg hover:bg-brand-200 dark:hover:bg-brand-600 transition-colors">{texto}</a>')
    partes = []
    for bloque in curso["bloques"]:
        detalles = []
        for l in bloque["lecciones"]:
            arch = archivo_base(l)
            cuerpo = "".join("<p>%s</p>" % escapar(p) for p in l["parrafos"])
            cuerpo += "<p><strong>Practica:</strong> %s</p>" % escapar(l["practica"])
            botones = (boton.format(slug=slug, arch=arch, suf=".pptx", texto="📊 Descargar presentación") +
                       boton.format(slug=slug, arch=arch, suf="-ejercicios.pdf", texto="📄 Descargar ejercicios (PDF)"))
            detalles.append(
                '<details class="bg-brand-50 dark:bg-brand-800 rounded-lg px-4 py-3">'
                '<summary class="cursor-pointer font-semibold text-brand-800 dark:text-white marker:text-accent-600">%d. %s</summary>'
                '<div class="mt-3 text-sm text-brand-600 dark:text-brand-300 space-y-2 leading-relaxed">%s</div>'
                '<div class="mt-3 flex flex-wrap gap-2">%s</div></details>'
                % (l["n"], escapar(l["titulo"]), cuerpo, botones))
        partes.append(
            '                    <div>\n'
            '                        <h4 class="font-serif text-base font-bold text-brand-800 dark:text-white mb-2">Bloque %d · %s</h4>\n'
            '                        <div class="space-y-2">\n                            %s\n'
            '                        </div>\n'
            '                    </div>' % (bloque["n"], escapar(bloque["titulo"]),
                                            "\n                            ".join(detalles)))
    return "\n".join(partes) + "\n"


# ------------------------------------------------------------- presentaciones
def presentacion(curso, leccion, destino):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN

    def rgb(h):
        return RGBColor.from_string(h)

    def caja(diapo, x, y, an, al, texto, tam, color, negrita=False, fuente="Calibri", alineado=None):
        tb = diapo.shapes.add_textbox(Inches(x), Inches(y), Inches(an), Inches(al))
        tf = tb.text_frame
        tf.word_wrap = True
        for i, linea in enumerate(texto.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            r = p.add_run()
            r.text = linea
            r.font.size = Pt(tam)
            r.font.bold = negrita
            r.font.name = fuente
            r.font.color.rgb = rgb(color)
            if alineado is not None:
                p.alignment = alineado
        return tb

    def figura(diapo, forma, x, y, an, al, color):
        f = diapo.shapes.add_shape(forma, Inches(x), Inches(y), Inches(an), Inches(al))
        f.fill.solid()
        f.fill.fore_color.rgb = rgb(color)
        f.line.fill.background()
        f.shadow.inherit = False
        return f

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    vacio = prs.slide_layouts[6]

    # 1 · portada de la lección
    d = prs.slides.add_slide(vacio)
    figura(d, MSO_SHAPE.RECTANGLE, 0, 0, 13.333, 7.5, AZUL_950)
    figura(d, MSO_SHAPE.OVAL, 10.6, -1.6, 5, 5, AZUL_900)
    figura(d, MSO_SHAPE.OVAL, -1.8, 5.4, 4.2, 4.2, AZUL_900)
    figura(d, MSO_SHAPE.OVAL, 0.7, 0.65, 1.1, 1.1, AMBAR_500)
    caja(d, 0.7, 0.65, 1.1, 1.1, str(leccion["n"]), 32, AZUL_950, True, alineado=PP_ALIGN.CENTER)
    caja(d, 2.1, 0.75, 10.4, 0.5, "AJEDREZ INTEGRAL  ·  %s" % curso["titulo"].upper(), 13, AMBAR_400, True)
    caja(d, 2.1, 1.2, 10.4, 0.4, "Bloque %d · %s" % (leccion["bloque"]["n"], leccion["bloque"]["titulo"]), 13, AZUL_300)
    caja(d, 0.7, 2.9, 11.9, 2.2, leccion["titulo"], 40, "FFFFFF", True, "Cambria")
    caja(d, 0.7, 5.05, 10.5, 1.0, leccion["resumen"], 18, CLARO)
    caja(d, 0.7, 6.9, 6.0, 0.4, "ajedrez-integral.com", 11, AZUL_400)

    # 2 · la idea (una diapositiva por párrafo, para que el texto respire)
    for parrafo in leccion["parrafos"]:
        d = prs.slides.add_slide(vacio)
        figura(d, MSO_SHAPE.OVAL, 0.6, 0.55, 0.55, 0.55, AMBAR_500)
        caja(d, 0.6, 0.55, 0.55, 0.55, "♟", 20, AZUL_950, alineado=PP_ALIGN.CENTER)
        caja(d, 1.35, 0.58, 8.0, 0.5, "CONCEPTO CLAVE", 14, AMBAR_500, True)
        caja(d, 0.6, 1.15, 12.1, 0.9, leccion["titulo"], 26, AZUL_950, True, "Cambria")
        figura(d, MSO_SHAPE.RECTANGLE, 0.6, 2.15, 12.1, 4.7, GRIS_50)
        caja(d, 1.1, 2.55, 11.2, 3.9, parrafo, 16, AZUL_600)

    # 3 · la tarea
    d = prs.slides.add_slide(vacio)
    figura(d, MSO_SHAPE.RECTANGLE, 0, 0, 13.333, 7.5, AZUL_900)
    figura(d, MSO_SHAPE.OVAL, -1.4, -1.6, 4.5, 4.5, AZUL_950)
    figura(d, MSO_SHAPE.OVAL, 5.66, 0.9, 2.0, 2.0, AMBAR_500)
    caja(d, 5.66, 0.9, 2.0, 2.0, "✔", 44, AZUL_950, True, alineado=PP_ALIGN.CENTER)
    caja(d, 0.6, 3.15, 12.1, 0.5, "PARA PRACTICAR", 14, AMBAR_400, True, alineado=PP_ALIGN.CENTER)
    caja(d, 1.4, 3.7, 10.5, 2.6, leccion["practica"], 22, "FFFFFF", False, "Cambria", PP_ALIGN.CENTER)
    caja(d, 0.6, 6.85, 12.1, 0.4, "%s · Lección %d" % (curso["titulo"], leccion["n"]), 12, AZUL_300,
         alineado=PP_ALIGN.CENTER)
    prs.save(destino)


# --------------------------------------------------------- PDF de ejercicios
def ejercicios(curso, leccion, destino):
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.lib.colors import HexColor
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable, PageBreak

    azul = HexColor("#" + AZUL_950)
    gris = HexColor("#" + AZUL_400)
    ambar = HexColor("#" + AMBAR_500)

    est = lambda **kw: ParagraphStyle(**kw)
    s_marca = est(name="marca", fontName="Helvetica-Bold", fontSize=9, textColor=ambar, spaceAfter=2, leading=12)
    s_sub = est(name="sub", fontName="Helvetica-Oblique", fontSize=9.5, textColor=gris, spaceAfter=10, leading=12)
    s_tit = est(name="tit", fontName="Times-Bold", fontSize=19, textColor=azul, spaceAfter=8, leading=23)
    s_sec = est(name="sec", fontName="Helvetica-Bold", fontSize=10, textColor=ambar, spaceBefore=12, spaceAfter=5, leading=13)
    s_txt = est(name="txt", fontName="Helvetica", fontSize=10.5, textColor=HexColor("#" + AZUL_600), leading=15, spaceAfter=7)
    s_item = est(name="item", fontName="Helvetica", fontSize=10.5, textColor=HexColor("#" + AZUL_600),
                 leading=15, spaceAfter=9, leftIndent=14, bulletIndent=2)
    s_ej = est(name="ej", fontName="Helvetica-Bold", fontSize=11, textColor=azul, spaceBefore=10, spaceAfter=3, leading=14)
    s_raya = est(name="raya", fontName="Helvetica", fontSize=10.5, textColor=gris, leading=26, spaceAfter=4)

    def pie(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(gris)
        canvas.drawString(2 * cm, 1.4 * cm, "Ajedrez Integral · ajedrez-integral.com")
        canvas.drawRightString(letter[0] - 2 * cm, 1.4 * cm, curso["titulo"])
        canvas.restoreState()

    doc = SimpleDocTemplate(destino, pagesize=letter, title="%s — Ejercicios %d" % (curso["titulo"], leccion["n"]),
                            author="Ajedrez Integral", leftMargin=2 * cm, rightMargin=2 * cm,
                            topMargin=2 * cm, bottomMargin=2.2 * cm)
    f = []
    f.append(Paragraph("AJEDREZ INTEGRAL &nbsp;·&nbsp; %s" % escapar(curso["titulo"]).upper(), s_marca))
    f.append(Paragraph("Ejercicios de práctica — Lección %d" % leccion["n"], s_sub))
    f.append(Paragraph(escapar(leccion["titulo"]), s_tit))
    f.append(HRFlowable(width="100%", thickness=1.2, color=ambar, spaceAfter=10))
    f.append(Paragraph("OBJETIVO DE LA LECCIÓN", s_sec))
    f.append(Paragraph(escapar(leccion["resumen"]), s_txt))
    f.append(Paragraph("REPASO RÁPIDO", s_sec))
    for p in leccion["parrafos"]:
        f.append(Paragraph(escapar(p), s_item, bulletText="–"))
    f.append(Paragraph("EJERCICIOS", s_sec))
    f.append(Paragraph("1. Práctica dirigida", s_ej))
    f.append(Paragraph(escapar(leccion["practica"]), s_txt))
    f.append(Paragraph("2. Explícalo con tus palabras", s_ej))
    f.append(Paragraph("Escribe en tres o cuatro líneas de qué se trata “%s” y para qué sirve, como si se lo "
                       "explicaras a alguien que recién empieza." % escapar(leccion["titulo"]), s_txt))
    for i in range(3):
        f.append(Paragraph("_" * 88, s_raya))
    f.append(PageBreak())
    f.append(Paragraph("AJEDREZ INTEGRAL &nbsp;·&nbsp; %s" % escapar(curso["titulo"]).upper(), s_marca))
    f.append(Paragraph("Ejercicios de práctica — Lección %d (continuación)" % leccion["n"], s_sub))
    f.append(HRFlowable(width="100%", thickness=1.2, color=ambar, spaceAfter=10))
    f.append(Paragraph("3. Pregunta de repaso", s_ej))
    f.append(Paragraph(escapar(leccion["parrafos"][0]), s_txt))
    f.append(Paragraph("¿Por qué es así? Escríbelo con tus palabras.", s_txt))
    for i in range(3):
        f.append(Paragraph("_" * 88, s_raya))
    f.append(Paragraph("4. Llévalo a una partida", s_ej))
    f.append(Paragraph("Juega una partida completa —contra Oscar en el Tablero de ajedrez-integral.com, contra el "
                       "motor o con un compañero— aplicando a propósito lo de esta lección. Al terminar, anota un "
                       "momento concreto en el que lo usaste, o en el que debiste usarlo.", s_txt))
    for i in range(5):
        f.append(Paragraph("_" * 88, s_raya))
    doc.build(f, onFirstPage=pie, onLaterPages=pie)


# ---------------------------------------------------------------------- main
def main(slug):
    ruta = os.path.join(RAIZ, "herramientas", "cursos", slug + ".json")
    curso = json.load(open(ruta, encoding="utf-8"))
    total = numerar(curso["bloques"])

    with open(os.path.join(RAIZ, "cursos", slug + ".html"), "w", encoding="utf-8") as fh:
        fh.write(portada(curso))
    with open(os.path.join(RAIZ, "cursos", "protegido", slug + ".html"), "w", encoding="utf-8") as fh:
        fh.write(protegido(curso))

    carpeta = os.path.join(RAIZ, "cursos", "recursos", slug)
    os.makedirs(carpeta, exist_ok=True)
    for bloque in curso["bloques"]:
        for leccion in bloque["lecciones"]:
            base = archivo_base(leccion)
            presentacion(curso, leccion, os.path.join(carpeta, base + ".pptx"))
            ejercicios(curso, leccion, os.path.join(carpeta, base + "-ejercicios.pdf"))

    print("Curso '%s': %d lecciones" % (curso["titulo"], total))
    print("  cursos/%s.html" % slug)
    print("  cursos/protegido/%s.html" % slug)
    print("  cursos/recursos/%s/ (%d archivos)" % (slug, total * 2))
    print("Falta agregar la tarjeta en cursos.html.")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "estrategia-en-el-final")
