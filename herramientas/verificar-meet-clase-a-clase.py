#!/usr/bin/env python3
"""Las grabaciones de Meet, clase a clase (herramientas/meet-clase-a-clase.py).

Con grabaciones inventadas y el motor falso (MEET_MOTOR=falso: el «audio» es
un texto), sin bajar Whisper ni salir a la red. Lo que se rompe acá se rompe
CALLADO —el script termina sin error y el informe queda con un hueco—, así
que se mira lo que deja:

- que cada grabación quieta deje su transcripción, su fila en las sesiones
  del JSON del mes y su entrada en clase-a-clase.md, en el mes que dice el
  nombre (no en el del día en que se corrió);
- que la que todavía se está bajando (cambió hace menos de dos minutos) y lo
  que no es audio ni video se dejen para después;
- que correrlo otra vez no repita filas ni pise lo corregido a mano en el
  JSON, y que una grabación nueva sí se sume;
- que la asistencia y los temas sugeridos vayan en amarillo («pend»), y que
  el Word del mes se arme y los traiga resaltados.

Ver «Las grabaciones de Meet, clase a clase» en docs/decisiones/informes.md.

Uso:  python3 herramientas/verificar-meet-clase-a-clase.py
"""
import json
import os
import re
import subprocess
import sys
import tempfile
import time

AQUI = os.path.dirname(os.path.abspath(__file__))
SCRIPT = os.path.join(AQUI, "meet-clase-a-clase.py")

fallos = 0


def ok(que, cumple, detalle=""):
    global fallos
    if cumple:
        print("  ✓", que)
    else:
        fallos += 1
        print("  ✗", que, ("\n      " + str(detalle)) if detalle else "")


def grabar(carpeta, nombre, lineas, hace_segundos=600):
    ruta = os.path.join(carpeta, nombre)
    with open(ruta, "w", encoding="utf-8") as f:
        f.write("\n".join(lineas))
    t = time.time() - hace_segundos
    os.utime(ruta, (t, t))


def correr(grab, trabajo):
    r = subprocess.run([sys.executable, SCRIPT, "--grabaciones", grab, "--trabajo", trabajo],
                       env={**os.environ, "MEET_MOTOR": "falso"}, capture_output=True, text=True)
    if r.returncode != 0:
        print(r.stdout, r.stderr)
    return r


def leer(ruta):
    with open(ruta, encoding="utf-8") as f:
        return json.load(f) if ruta.endswith(".json") else f.read()


print("\n=== Las grabaciones de Meet, clase a clase ===")
tmp = tempfile.mkdtemp(prefix="meet-")
grab, trabajo = os.path.join(tmp, "Meet Recordings"), os.path.join(tmp, "Informe CCDR")
os.makedirs(grab)
os.makedirs(trabajo)
with open(os.path.join(trabajo, "estudiantes.csv"), "w", encoding="utf-8-sig") as f:
    f.write("Estudiante;Ejercicios\nAtleta Prueba Uno;10\nAtleta Prueba Dos;4\n")

clavada = ["Hoy vemos la clavada.", "Esta pieza está clavada contra el rey.", "Otra clavada en la diagonal.",
           "Romper la clavada.", "Repasamos."]
finales = ["Finales de rey y peón.", "La oposición es clave en el final.", "Otra vez la oposición.",
           "Regla del cuadrado.", "El rey activo en el final.", "Lucena."]
grabar(grab, "Clase CCDR (2030-01-28 18:02 GMT-6).mp4", clavada)
grabar(grab, "Clase CCDR - 2030_01_07 18:00 CST - Recording.webm", finales)
grabar(grab, "Clase CCDR (2030-02-04 18:00 GMT-6).mp4", ["Una clase de febrero."])
grabar(grab, "Bajándose (2030-01-30 18:00 GMT-6).mp4", ["a medias"], hace_segundos=5)
grabar(grab, "Clase CCDR (2030-01-28 18:02 GMT-6) - Transcript.txt", ["no es audio"])

r = correr(grab, trabajo)
ok("corre sin error", r.returncode == 0, r.stderr[-400:])

ene = os.path.join(trabajo, "2030-01")
datos = leer(os.path.join(ene, "datos-2030-01.json"))
ses = datos["lugar"]["sesiones"]
ok("dos sesiones de enero, en orden de fecha (aunque los nombres digan la fecha distinto)",
   [s[0] for s in ses] == ["Lun 07/01", "Lun 28/01"], [s[0] for s in ses])
ok("el horario sale del nombre y de la duración", ses[1][2] == "6:02 p. m. – 6:07 p. m.", ses[1][2])
ok("la sesión lleva el nombre de la reunión y los temas en amarillo",
   ses[1][3] == ["Clase CCDR", {"pend": " — temas: Táctica: clavada"}], ses[1][3])
ok("la asistencia va en amarillo, sin inventar un número", ses[0][4] == {"pend": "[completar]"}, ses[0][4])
ok("los finales se reconocen", "Finales" in json.dumps(ses[0][3], ensure_ascii=False), ses[0][3])
ok("los atletas salen de estudiantes.csv", datos["atletas"] == ["Atleta Prueba Uno", "Atleta Prueba Dos"], datos["atletas"])
ok("el periodo es el mes de la grabación", datos["periodo"] == "1 al 31 de enero de 2030", datos["periodo"])

txts = sorted(os.listdir(os.path.join(ene, "transcripciones")))
ok("una transcripción por grabación de enero", txts == ["2030-01-07 18.00 Clase CCDR.txt", "2030-01-28 18.02 Clase CCDR.txt"], txts)
t = leer(os.path.join(ene, "transcripciones", txts[1]))
ok("la transcripción trae las marcas de tiempo y el aviso de que es automática",
   "[00:01:00] Esta pieza está clavada contra el rey." in t and "automática" in t, t[:300])

ok("la de febrero va a su propio mes", os.path.exists(os.path.join(trabajo, "2030-02", "datos-2030-02.json")))
hechas = leer(os.path.join(trabajo, "procesadas.json"))
ok("la que se está bajando y lo que no es audio quedan para después",
   not any(n.startswith("Bajándose") or n.endswith(".txt") for n in hechas), list(hechas))

md = leer(os.path.join(ene, "clase-a-clase.md"))
ok("clase-a-clase.md suma el mes", "2 sesiones grabadas en Meet" in md and "## Lun 28/01" in md, md[:300])

# Lo corregido a mano se queda, y nada se repite.
ses[0][4] = "9"
datos["otros"] = [{"p": "Escrito a mano."}]
with open(os.path.join(ene, "datos-2030-01.json"), "w", encoding="utf-8") as f:
    json.dump(datos, f, ensure_ascii=False)
# La nueva, dentro de una subcarpeta (quien las ordena por grupo o por mes).
os.makedirs(os.path.join(grab, "Ordenadas", "Enero"))
grabar(os.path.join(grab, "Ordenadas", "Enero"), "Clase CCDR (2030-01-14 18:00 GMT-6).mp4", ["Mate en dos.", "Otro mate en dos.", "Jaque mate."])
# Y una ya hecha, movida a otra subcarpeta: no se vuelve a transcribir.
os.rename(os.path.join(grab, "Clase CCDR (2030-01-28 18:02 GMT-6).mp4"),
          os.path.join(grab, "Ordenadas", "Clase CCDR (2030-01-28 18:02 GMT-6).mp4"))
antes = os.path.getmtime(os.path.join(ene, "transcripciones", "2030-01-28 18.02 Clase CCDR.txt"))
correr(grab, trabajo)
datos = leer(os.path.join(ene, "datos-2030-01.json"))
ses = datos["lugar"]["sesiones"]
ok("otra vuelta suma la nueva (de una subcarpeta) en su lugar y no repite las de antes",
   [s[0] for s in ses] == ["Lun 07/01", "Lun 14/01", "Lun 28/01"], [s[0] for s in ses])
ok("la que se movió a una subcarpeta no se vuelve a transcribir",
   os.path.getmtime(os.path.join(ene, "transcripciones", "2030-01-28 18.02 Clase CCDR.txt")) == antes)
ok("lo corregido a mano se queda", ses[0][4] == "9" and datos["otros"] == [{"p": "Escrito a mano."}], (ses[0][4], datos["otros"]))

# El Word del mes, con lo pendiente resaltado.
docx = os.path.join(ene, "Informe-en-curso-CCDR-enero-2030.docx")
if not os.path.exists(docx):
    ok("se arma el Word del mes (¿falta `npm install`?)", False)
else:
    try:
        from docx import Document
    except ImportError:
        sys.exit("Falta python-docx:  pip install python-docx")
    d = Document(docx)
    # La tabla de sesiones va DENTRO del cuadro de la pregunta: hay que bajar.
    def todas(tablas):
        for tb in tablas:
            for row in tb.rows:
                for c in row.cells:
                    yield c
                    yield from todas(c.tables)
    celdas = list(todas(d.tables))
    sesion = next((c for c in celdas if c.text.startswith("Clase CCDR — temas: Mates")), None)
    ok("el Word trae la sesión nueva", sesion is not None)
    amarillo = [r.text for c in celdas for p in c.paragraphs for r in p.runs if r.font.highlight_color is not None]
    ok("en el Word, los temas y la asistencia sin contar salen en amarillo",
       " — temas: Mates" in amarillo and "[completar]" in amarillo, amarillo[:6])

# --registro: con pythonw no hay ventana, y esto es lo único que queda.
registro = os.path.join(tmp, "registro.txt")
r = subprocess.run([sys.executable, SCRIPT, "--grabaciones", grab, "--trabajo", trabajo, "--registro", registro],
                   env={**os.environ, "MEET_MOTOR": "falso"}, capture_output=True, text=True)
ok("con --registro, lo que pasa queda en el archivo y no en la consola",
   r.returncode == 0 and "Arranca:" in leer(registro) and not r.stdout, (r.returncode, r.stdout[:200]))

# El instalador de Windows (aquí no hay PowerShell para correrlo): que pase
# solo opciones que el script conoce, y que lleve el BOM sin el que
# PowerShell 5 rompe las tildes.
with open(os.path.join(AQUI, "meet-clase-a-clase-instalar.ps1"), "rb") as f:
    ps1 = f.read()
ok("el instalador de Windows va con BOM", ps1.startswith(b"\xef\xbb\xbf"))
ps1 = ps1.decode("utf-8-sig")
argumentos = next((l for l in ps1.splitlines() if l.strip().startswith("$argumentos =")), "")
with open(SCRIPT, encoding="utf-8") as f:
    conocidas = set(re.findall(r'add_argument\("(--[a-z-]+)"', f.read()))
pasa = set(re.findall(r"(--[a-z-]+)", argumentos))
ok("el instalador pasa al script solo opciones que existen, con --vigilar y --registro",
   pasa and pasa <= conocidas and {"--vigilar", "--registro"} <= pasa, (pasa - conocidas, argumentos))

print("\n" + ("Todo bien." if not fallos else f"{fallos} fallo(s)."))
sys.exit(1 if fallos else 0)
