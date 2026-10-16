#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Arma el PowerPoint de una sesión de Formación Ajedrez con sus láminas.

    python3 herramientas/formacion-pptx.py clase-01 [clase-02 …]
    python3 herramientas/formacion-pptx.py            (todas las que tienen láminas)

    pip install python-pptx pillow

El .pptx que se baja desde el curso tiene que ser lo mismo que se proyecta en
la clase en vivo. Antes salía de otro lado (el texto completo del JSON del
curso, en el formato de antes de «Las láminas dicen poco») y se quedó viejo:
el lunes 12 de octubre la Clase 1 se dio con 51 láminas nuevas y el .pptx que
se descargaba era el del domingo. Ahora sale de las mismas láminas que dibuja
herramientas/formacion-laminas.js (cursos/recursos/formacion-ajedrez/
presentaciones/clase-NN/): una imagen por diapositiva, a toda la diapositiva,
con las notas de diapositivas.json en las notas del orador (las ve el profe al
presentar) y su `texto` como texto alternativo de la imagen (lo lee el lector
de pantalla).

En las propiedades del archivo (Asunto) queda la huella de lo que se usó: el
sha-256 de las imágenes y de diapositivas.json. verificar-formacion-pptx.py la
vuelve a calcular: si las láminas cambiaron y el .pptx no, falla.

Lo llama herramientas/curso-generar-formacion.py después de dibujar las
láminas; para volver a armar solo el PowerPoint, alcanza con este script.
"""
import datetime
import hashlib
import io
import json
import os
import re
import sys
import zipfile

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RECURSOS = os.path.join(RAIZ, "cursos", "recursos", "formacion-ajedrez")
PRESENTACIONES = os.path.join(RECURSOS, "presentaciones")
DATOS_CURSO = os.path.join(RAIZ, "herramientas", "cursos", "formacion-ajedrez.json")
DATOS_LAMINAS = os.path.join(RAIZ, "herramientas", "datos")


def huella(clase):
    """El sha-256 de diapositivas.json y de cada imagen, en orden. Lo mismo
    calcula verificar-formacion-pptx.py."""
    carpeta = os.path.join(PRESENTACIONES, clase)
    with open(os.path.join(carpeta, "diapositivas.json"), "rb") as fh:
        crudo = fh.read()
    h = hashlib.sha256(crudo)
    for d in json.loads(crudo)["diapositivas"]:
        with open(os.path.join(carpeta, d["imagen"]), "rb") as fh:
            h.update(fh.read())
    return h.hexdigest()


def archivo_de(clase):
    """El nombre del .pptx de la sesión NN, del JSON del curso (las sesiones
    se numeran en orden, como numerar() del generador)."""
    n = int(clase.split("-")[1])
    curso = json.load(open(DATOS_CURSO, encoding="utf-8"))
    lecciones = [l for b in curso["bloques"] for l in b["lecciones"]]
    return lecciones[n - 1]["archivo"]


def clases_con_laminas():
    return sorted(m.group(1) for m in (re.match(r"^formacion-laminas-(clase-\d+)\.js$", f)
                                       for f in os.listdir(DATOS_LAMINAS)) if m)


def armar(clase):
    from pptx import Presentation
    from pptx.util import Emu
    from PIL import Image

    carpeta = os.path.join(PRESENTACIONES, clase)
    datos = json.load(open(os.path.join(carpeta, "diapositivas.json"), encoding="utf-8"))
    destino = os.path.join(RECURSOS, archivo_de(clase) + ".pptx")

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)   # 16:9, como las láminas
    vacia = prs.slide_layouts[6]
    for i, d in enumerate(datos["diapositivas"], 1):
        # PowerPoint viejo y LibreOffice no abren WebP dentro de un .pptx: va en
        # JPEG (en PNG, la Clase 1 pesaba 13 MB; así, la cuarta parte).
        img = Image.open(os.path.join(carpeta, d["imagen"])).convert("RGB")
        jpg = io.BytesIO()
        img.save(jpg, "JPEG", quality=90, optimize=True)
        jpg.seek(0)
        diapo = prs.slides.add_slide(vacia)
        foto = diapo.shapes.add_picture(jpg, 0, 0, prs.slide_width, prs.slide_height)
        foto.name = "Lámina %d" % i
        foto._element.nvPicPr.cNvPr.set("descr", d.get("texto") or d["titulo"])
        if d.get("notas"):
            diapo.notes_slide.notes_text_frame.text = d["notas"]

    prs.core_properties.title = datos["titulo"]
    prs.core_properties.author = "Ajedrez Integral"
    prs.core_properties.subject = "laminas:" + huella(clase)
    prs.core_properties.keywords = ""
    prs.core_properties.comments = ""
    prs.core_properties.last_modified_by = "Ajedrez Integral"
    prs.core_properties.revision = 1
    # Fechas fijas, también las del zip: el mismo contenido da el mismo archivo
    # y git no ve cambios de más.
    fija = datetime.datetime(2026, 1, 1)
    prs.core_properties.created = fija
    prs.core_properties.modified = fija
    crudo = io.BytesIO()
    prs.save(crudo)
    with zipfile.ZipFile(crudo) as ent, zipfile.ZipFile(destino, "w", zipfile.ZIP_DEFLATED) as sal:
        for info in ent.infolist():
            nueva = zipfile.ZipInfo(info.filename, date_time=fija.timetuple()[:6])
            nueva.compress_type = zipfile.ZIP_DEFLATED
            sal.writestr(nueva, ent.read(info.filename))
    print("  %s: %d diapositivas → %s" % (clase, len(datos["diapositivas"]), os.path.relpath(destino, RAIZ)))


def main():
    clases = sys.argv[1:] or clases_con_laminas()
    for clase in clases:
        if not re.match(r"^clase-\d+$", clase):
            raise SystemExit("Uso: formacion-pptx.py [clase-NN …]")
        armar(clase)


if __name__ == "__main__":
    main()
