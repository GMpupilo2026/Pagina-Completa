#!/usr/bin/env python3
"""Revisa el Word y el PDF que dejó herramientas/verificar-reportes.js.

Va aparte del script de navegador a propósito: abrir un .docx y un .pdf pide
librerías de Python (pypdf, python-docx) que no tienen nada que hacer dentro de
una prueba de navegador, y acá se comprueba lo que de verdad importa y no se ve
en pantalla — que los dos archivos ABRAN, que lleven la marca de agua en TODAS
las páginas y que el texto haya llegado con sus tildes.

Uso:  node herramientas/verificar-reportes.js         (deja los archivos)
      python3 herramientas/verificar-reportes.py <carpeta>
"""
import sys, glob, os, re, zipfile

fallos = 0

def cumple(nombre, condicion, detalle=""):
    global fallos
    if condicion:
        print("  ✓ " + nombre + ((": " + str(detalle)) if detalle else ""))
    else:
        print("  ✗ " + nombre + ((" — " + str(detalle)) if detalle else ""))
        fallos += 1


def revisar_pdf(ruta):
    from pypdf import PdfReader
    print("\n=== El PDF ===")
    lector = PdfReader(ruta)
    paginas = len(lector.pages)
    cumple("abre y tiene páginas", paginas > 0, str(paginas) + " páginas")
    cumple("lleva el autor en los datos del archivo",
           (lector.metadata or {}).get("/Author"), (lector.metadata or {}).get("/Author"))

    # La marca de agua tiene que estar en TODAS, no solo en la primera: es el
    # error clásico de estampar solo la portada.
    sin_marca = [i + 1 for i, p in enumerate(lector.pages)
                 if "/Marca" not in (p.get("/Resources", {}).get("/XObject", {}) or {})]
    cumple("la marca de agua está en todas las páginas", not sin_marca,
           "faltan en " + str(sin_marca) if sin_marca else str(paginas) + " de " + str(paginas))

    texto = "\n".join((p.extract_text() or "") for p in lector.pages)
    cumple("se lee el título", "Informe de actividades de clases" in texto)
    cumple("trae las fechas de las clases", "14/09/2026" in texto and "15/09/2026" in texto)
    cumple("trae la asistencia", "Jean Quesada Arauz" in texto)
    cumple("las tildes y la eñe llegaron",
           all(p in texto for p in ["Táctica", "María", "Duración", "niño"]))
    cumple("el guion largo no se perdió", "—" in texto)
    cumple("lleva el pie con el autor en todas",
           all("Oscar Angulo Cubero" in (p.extract_text() or "") for p in lector.pages))
    return paginas


def revisar_docx(ruta):
    print("\n=== El Word ===")
    z = zipfile.ZipFile(ruta)
    cumple("el ZIP está sano", z.testzip() is None)
    cumple("[Content_Types].xml va primero, que es lo que pide el formato",
           z.namelist()[0] == "[Content_Types].xml", z.namelist()[0])

    import xml.dom.minidom as md
    for parte in [n for n in z.namelist() if n.endswith(".xml") or n.endswith(".rels")]:
        try:
            md.parseString(z.read(parte))
        except Exception as e:
            cumple("XML bien formado: " + parte, False, str(e))
            return
    cumple("todos los XML están bien formados", True, str(len(z.namelist())) + " partes")

    try:
        import docx
        d = docx.Document(ruta)
        cumple("un lector de Word de verdad lo abre", True,
               str(len(d.paragraphs)) + " párrafos, " + str(len(d.tables)) + " tablas")
        texto = "\n".join(p.text for p in d.paragraphs)
        for t in d.tables:
            for fila in t.rows:
                texto += "\n" + " ".join(c.text for c in fila.cells)
    except ImportError:
        texto = "\n".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>",
                                     z.read("word/document.xml").decode("utf-8")))
        cumple("un lector de Word de verdad lo abre", False, "falta python-docx, se revisó el XML a mano")

    cumple("se lee el título", "Informe de actividades de clases" in texto)
    cumple("trae las fechas de las clases", "14/09/2026" in texto)
    cumple("trae la asistencia", "Jean Quesada Arauz" in texto)
    cumple("las tildes y la eñe llegaron",
           all(p in texto for p in ["Táctica", "María", "Duración"]))

    encabezado = z.read("word/header1.xml").decode("utf-8")
    cumple("la marca de agua va en el encabezado, o sea en todas las páginas",
           "rIdMarca" in encabezado and "MarcaDeAgua" in encabezado)
    cumple("y la imagen de la marca está dentro del archivo",
           "word/media/marca.png" in z.namelist())
    cumple("el pie lleva el nombre del autor",
           "Oscar Angulo Cubero" in z.read("word/footer1.xml").decode("utf-8"))


def revisar_externo(ruta):
    """El informe de clases dadas por fuera: los números salen de las hojas."""
    from pypdf import PdfReader
    print("\n=== El PDF del informe de clases externas ===")
    lector = PdfReader(ruta)
    cumple("abre y tiene páginas", len(lector.pages) > 0, str(len(lector.pages)) + " páginas")
    sin_marca = [i + 1 for i, p in enumerate(lector.pages)
                 if "/Marca" not in (p.get("/Resources", {}).get("/XObject", {}) or {})]
    cumple("la marca de agua está en todas las páginas", not sin_marca,
           "faltan en " + str(sin_marca) if sin_marca else "todas")
    texto = "\n".join((p.extract_text() or "") for p in lector.pages)
    cumple("trae la asistencia por clase y por estudiante",
           "Asistencia por clase" in texto and "Asistencia por estudiante" in texto)
    cumple("dice quién faltó", "No vinieron" in texto)
    cumple("trae el contenido que venía en el documento",
           "Empezamos con el caballo" in texto)
    # Que el informe EXPLIQUE cómo leyó la hoja no es un adorno: una cuadrícula
    # mal entendida da números creíbles y falsos, y esta es la única forma de
    # que el error se vea desde el papel.
    cumple("y explica cómo leyó la hoja de asistencia",
           "De dónde salen estos números" in texto and "en blanco se contaron como falta" in texto)


if __name__ == "__main__":
    carpeta = sys.argv[1] if len(sys.argv) > 1 else ""
    if not carpeta or not os.path.isdir(carpeta):
        print("Uso: python3 herramientas/verificar-reportes.py <carpeta que dejó el script de node>")
        sys.exit(2)
    # El script de navegador deja también el PDF del informe de clases externas,
    # con el prefijo "externo-". Dice otras cosas, así que se revisa aparte: si
    # se mezclaran, las comprobaciones fallarían por mirar el archivo que no es.
    todos = glob.glob(os.path.join(carpeta, "*.pdf"))
    pdfs = [f for f in todos if not os.path.basename(f).startswith("externo-")]
    externos = [f for f in todos if os.path.basename(f).startswith("externo-")]
    docxs = glob.glob(os.path.join(carpeta, "*.docx"))
    if not pdfs or not docxs:
        print("En esa carpeta no están los dos archivos.")
        sys.exit(2)
    revisar_pdf(pdfs[0])
    revisar_docx(docxs[0])
    if externos:
        revisar_externo(externos[0])
    print("\n" + (str(fallos) + " fallo(s)" if fallos else "Los dos archivos, bien."))
    sys.exit(1 if fallos else 0)
